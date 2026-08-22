"""
Assemble the TRAP region-selection + zoom-in figures into one presentation.

Story flow (why cluster first, then zoom in):
  Part 0  Roadmap / rationale (funnel: hundreds of regions -> clusters -> 7 regions -> zoom-in)
  Part 1  How the 7 regions were selected      (Fig01..Fig07, region_selection_PPT/)
  Part 2  Zoom-in per region, group level      (<REG>_density_count_by_phase.png)
  Part 3  Zoom-in per region, within-phase      (<REG>_withinphase_ActiveVsPassive.png)

Each slide carries a short, plain-language explanation of WHAT it shows and WHY.

Run:  python trap_build_pptx.py
Out:  TRAP_OUTPUT_calculated_mm3/TRAP_region_selection_and_zoomin.pptx
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

ROOT = Path(r"C:\Users\hsollim\behavior_task\TRAP_analysis_sync")
OUTBASE = ROOT / "TRAP_OUTPUT_calculated_mm3"
SEL = OUTBASE / "region_selection_PPT"
ZOOM = OUTBASE / "region_zoomin"
SCAN = OUTBASE / "withdrawal_dip_scan"
PPTX = OUTBASE / "TRAP_region_selection_and_zoomin.pptx"

REGIONS = ["ORBm", "CA", "AId", "BMAp", "LM", "RE", "CP"]
REGION_CLUSTER = {"ORBm": 1, "CA": 1, "AId": 1, "BMAp": 4, "LM": 4, "RE": 4, "CP": 4}
REGION_INFO = {
    "CP":   ("Caudoputamen (striatum)", "goal-directed → habitual/compulsive drug seeking"),
    "BMAp": ("Basomedial amygdala, post.", "emotional valence, reward/threat state, addiction motivation"),
    "LM":   ("Lateral mammillary n. (hypothalamus)", "contextual/spatial memory of drug seeking"),
    "RE":   ("Nucleus reuniens (midline thalamus)", "mPFC–hippocampus hub; relapse & internal-state signalling"),
    "ORBm": ("Orbitofrontal cortex, medial", "outcome-value evaluation, on-task behaviour"),
    "CA":   ("Hippocampus CA", "context & associative memory of drug experience"),
    "AId":  ("Agranular insula, dorsal", "interoception, aversive/pain state, salience"),
}

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xB3, 0x00, 0x00)
BLUE = RGBColor(0x0A, 0x2F, 0x6B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def _title(slide, text, color=NAVY, size=26):
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(SW - 0.9), Inches(0.85))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(size); p.font.bold = True; p.font.color.rgb = color
    return tb


def _img_fit(slide, img_path, left, top, max_w, max_h):
    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / h
    width = max_w
    height = width / ar
    if height > max_h:
        height = max_h
        width = height * ar
    left2 = left + (max_w - width) / 2
    top2 = top + (max_h - height) / 2
    slide.shapes.add_picture(str(img_path), Inches(left2), Inches(top2), Inches(width), Inches(height))


def _bullets(slide, left, top, width, height, items, size=14, header=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    if header:
        p = tf.paragraphs[0]; p.text = header
        p.font.size = Pt(size + 2); p.font.bold = True; p.font.color.rgb = NAVY
        first = False
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = it
        p.font.size = Pt(size); p.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        p.space_after = Pt(6)
    return tb


def slide_text(title, blocks, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    _title(s, title)
    if subtitle:
        tb = s.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(SW - 1), Inches(0.5))
        p = tb.text_frame.paragraphs[0]; p.text = subtitle
        p.font.size = Pt(16); p.font.italic = True; p.font.color.rgb = GREY
    _bullets(s, 0.7, 1.7, SW - 1.4, SH - 2.2, blocks, size=18)
    return s


def slide_image_right_text(title, img_path, bullets, header=None):
    s = prs.slides.add_slide(BLANK)
    _title(s, title)
    _img_fit(s, img_path, left=0.35, top=1.15, max_w=8.7, max_h=SH - 1.5)
    _bullets(s, 9.2, 1.3, 3.9, SH - 1.8, bullets, size=14, header=header)
    return s


def slide_image_big(title, img_path, caption):
    s = prs.slides.add_slide(BLANK)
    _title(s, title)
    _img_fit(s, img_path, left=0.5, top=1.15, max_w=SW - 1.0, max_h=SH - 2.1)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(SH - 0.85), Inches(SW - 1.2), Inches(0.7))
    p = tb.text_frame.paragraphs[0]; p.text = caption
    p.font.size = Pt(13); p.font.color.rgb = GREY; p.alignment = PP_ALIGN.CENTER
    tb.text_frame.word_wrap = True
    return s


# ---------------------------------------------------------------- title
s = prs.slides.add_slide(BLANK)
box = s.shapes.add_textbox(Inches(1), Inches(2.4), Inches(SW - 2), Inches(2))
tf = box.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = "Whole-brain TRAP → 7 candidate regions for Xenium"
p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = NAVY
p2 = tf.add_paragraph(); p2.text = "Unbiased clustering to filter hundreds of regions, then per-region zoom-in"
p2.font.size = Pt(20); p2.font.color.rgb = GREY
p3 = tf.add_paragraph(); p3.text = "Active (morphine) vs Passive (yoked) · phases: During → Post → Withdrawal → Reinstatement"
p3.font.size = Pt(15); p3.font.italic = True; p3.font.color.rgb = GREY

# ---------------------------------------------------------------- roadmap / why
slide_text(
    "Roadmap — why cluster first, then zoom in",
    [
        "Problem: the brain has hundreds of regions. Testing each one individually is under-powered "
        "and hard to interpret (multiple-comparisons, no structure).",
        "Step 1 — FILTER by clustering: pool all 138 forebrain regions (L/R-averaged; no brainstem/"
        "cerebellum/white-matter) and group them by their Active-vs-Passive activity fingerprint across "
        "all 4 phases (universal PCA + k-means, K=4). Hundreds of regions → 4 functional clusters.",
        "Step 2 — KEEP behaviourally-aligned clusters: only clusters 1 & 4 reproduce the motivation "
        "timeline (Active > Passive at Post & Reinstatement). → shortlist of candidate regions.",
        "Step 3 — SELECT anatomically-relevant regions: BMAp, LM, RE, CP (cluster 4) + ORBm, CA, AId "
        "(cluster 1) = 7 regions, agreeing with an independent top-N ranking.",
        "Step 4 — ZOOM IN: for each of the 7 regions, look at density AND cell count across phases, "
        "Active vs Passive, showing every individual mouse.",
    ],
    subtitle="Funnel logic: hundreds of regions → 4 clusters → 2 behaviour-relevant clusters → 7 regions → zoom-in",
)

# ---------------------------------------------------------------- Part 1 header
slide_text("Part 1 · How we selected the 7 regions (clustering)",
           ["Figures 1–7 show the method, what the clusters look like, why clusters 1 & 4 were chosen, "
            "and why these specific 7 regions — all from the real Step-3/Step-13 clustering output."])

fig_desc = {
    "Fig01_clustering_method_PCAmap.png": (
        "Fig 1 · How clustering was done + what clusters look like",
        "Method (Fig 1)",
        ["A: 138 forebrain regions in PCA space, colored by cluster (K=4); circled = 7 selected.",
         "B: k was chosen at 4 (balances interpretability, variance, behaviour mapping).",
         "C: each region z-scored across samples (4 phases pooled) → k-means (sq-Euclidean, 50 restarts).",
         "PCA is only for 2-D display; clusters use the full profile."]),
    "Fig02_cluster_trajectories_AminusP.png": (
        "Fig 2 · Cluster trajectories: Active − Passive across phases",
        "Trajectories (Fig 2)",
        ["Each line = one cluster's mean (Active − Passive) across During→Post→Withdrawal→Reinstatement.",
         "Clusters 1 & 4 rise above 0 (Active higher) at Post & Reinstatement and dip at Withdrawal.",
         "This mirrors the motivational-behaviour timeline."]),
    "Fig03_why_clusters_1_and_4.png": (
        "Fig 3 · Why clusters 1 and 4 were selected",
        "Selection of clusters (Fig 3)",
        ["Per-phase bars of cluster-mean Active − Passive.",
         "Only clusters 1 & 4 show the predicted pattern (A>P at Post & Reinstatement).",
         "Post = Active morphine learning; Withdrawal = Passive sucrose; Reinstatement = Active craving.",
         "Clusters 2 & 3 lack this signature → not selected."]),
    "Fig04_cluster4_region_selection.png": (
        "Fig 4 · Cluster 4 — selection decided by the Withdrawal phase",
        "Within cluster 4 (Fig 4)",
        ["23 cluster-4 regions × phase heatmap of Active − Passive.",
         "SOLID black box = SELECTED (incubation): Active>Passive in craving phases AND still A>P in "
         "Withdrawal → BMAp, LM, RE, CP.",
         "DASHED green box = behaviour-match alternative: A>P craving but A<P in Withdrawal → PVi, SI, TTv, OT.",
         "We kept the incubation set on purpose (see next slide)."]),
    "Fig05_cluster1_region_selection.png": (
        "Fig 5 · Cluster 1 — selection decided by the Withdrawal phase",
        "Within cluster 1 (Fig 5)",
        ["Top cluster-1 regions by Post+Reinstatement Active − Passive.",
         "SOLID = SELECTED (incubation): A>P persists in Withdrawal → ORBm, CA, AId.",
         "DASHED green = behaviour-match: A>P craving but A<P in Withdrawal → ENT, ProS, EPd, IntG.",
         "Both are defensible; the difference is the Withdrawal phase."]),
    "Fig06_7region_crosscheck.png": (
        "Fig 6 · Cross-check — the 7 regions are consistently Active-dominant",
        "Cross-check (Fig 6)",
        ["All 7 regions show Active > Passive in the craving-relevant phases (Post & Reinstatement).",
         "Agrees with an independent top-N ranking (not just clustering).",
         "Gives confidence the selection is robust."]),
    "Fig07_final_7region_summary.png": (
        "Fig 7 · Final 7 regions for Xenium",
        "Summary (Fig 7)",
        ["Cluster 4: BMAp, LM, RE, CP (amygdala/striatal/thalamic/hypothalamic).",
         "Cluster 1: ORBm, CA, AId (frontal/hippocampal/insular).",
         "Each region's addiction relevance is listed for the Xenium rationale."]),
}
for fname, (title, header, bullets) in fig_desc.items():
    if fname == "Fig07_final_7region_summary.png":
        continue  # rendered after the two-philosophy figure
    slide_image_right_text(title, SEL / fname, bullets, header=header)

# ---- two selection philosophies (Withdrawal phase) ----
slide_text(
    "Selection nuance — the Withdrawal phase decides which regions we keep",
    [
        "Behaviour in Withdrawal: the Passive group shifts toward sucrose (natural reward), so behaviourally "
        "Passive > Active in Withdrawal.",
        "Philosophy A — INCUBATION of craving (what we chose): we KEPT regions where Active is STILL > Passive "
        "during Withdrawal. Rationale: persistent Active-dominant activity through withdrawal is what could "
        "drive the gradual incubation of craving. → BMAp, LM, RE, CP, ORBm, CA, AId.",
        "Philosophy B — EXACT behaviour match (alternative): if instead we want regions that track the behaviour "
        "one-to-one (Active < Passive in Withdrawal, like the sucrose shift), the candidates are PVi, SI, ENT "
        "(and ProS, EPd, IntG, TTv, OT).",
        "Both are scientifically valid; they answer different questions. We prioritised the incubation hypothesis, "
        "but the behaviour-match regions remain strong alternatives to discuss.",
    ],
    subtitle="Why the 7 selected regions stay Active>Passive in Withdrawal — and what the alternative would be",
)
slide_image_big(
    "Fig 8 · Two selection philosophies — incubation (kept) vs behaviour-match (alternative)",
    SEL / "Fig08_two_selection_philosophies.png",
    "Top (red) block = selected incubation regions (Active>Passive even in Withdrawal). "
    "Bottom (green) block = behaviour-match alternatives (Active<Passive in Withdrawal) e.g. PVi, SI, ENT.")

slide_image_right_text(*[
    "Fig 7 · Final 7 regions for Xenium",
    SEL / "Fig07_final_7region_summary.png",
    fig_desc["Fig07_final_7region_summary.png"][2],
], header=fig_desc["Fig07_final_7region_summary.png"][1])

# ---- Active-SELECTIVE withdrawal-dip signature (per-group trajectories) ----
slide_text(
    "Is the effect Active-SELECTIVE? — the ORBm / BMAp signature",
    [
        "Key question for a clean story: does the c-Fos change belong to the ACTIVE (morphine) group only, "
        "or do both groups move together? We compared each group's OWN density trajectory across phases.",
        "ORBm & BMAp: Active c-Fos RISES at Post, DIPS at Withdrawal, partially rebounds at Reinstatement — and "
        "the PASSIVE group shows essentially NO change. This Active-selective pattern is exactly what an "
        "addiction/craving ensemble should look like → strong rationale for taking ORBm & BMAp forward.",
        "AId: BOTH groups show the SAME rise-and-dip (Active ≈ Passive), so AId does not isolate the "
        "active/craving process — it is the weakest of the selected regions on this criterion.",
        "CA: the Active trajectory is PASSIVE-like (no clean craving peak); CA tracks the passive/behaviour "
        "pattern more than the active-craving one.",
        "Conclusion: on the Active-selectivity criterion, ORBm and BMAp are the most defensible next targets.",
    ],
    subtitle="Per-group density trajectories: Active-only rise+dip vs both-groups-same vs passive-like",
)
slide_image_big(
    "Fig 9b · Per-group density trajectories (z across phases) — Active vs Passive",
    SCAN / "Fig09b_group_trajectories_top.png",
    "ORBm/BMAp = Active-selective (red rises at Post, dips at Withdrawal; blue Passive flat). "
    "AId = both groups overlap. CA = Active is passive-like. GPi/MA = other regions with the same signature. "
    "SI/IntG/PVi = same shape but a DEEPER dip (Active falls below Passive).")
slide_text(
    "Re-check — do OTHER cluster-1/4 regions show the same ORBm/BMAp signature?",
    [
        "We scanned ALL 103 cluster-1 & 4 regions for the same signature (Active: Post peak + Withdrawal dip, "
        "Passive flat), then split them by the Withdrawal Active−Passive sign.",
        "ORBm/BMAp-TYPE (dip but Active still ≥ Passive in Withdrawal — incubation-compatible): besides the "
        "already-selected ORBm, BMAp, LM, RE, the only additional regions are GPi, VAL, PAR, APr, MA — all "
        "non-canonical for addiction, so ORBm/BMAp remain the best-justified choices.",
        "SAME shape but DEEPER dip (Active falls BELOW Passive in Withdrawal → behaviour-match, not incubation): "
        "a large set incl. PVi, SI, ENT, IntG, EPd, ProS, MS, TTv, OT — these track behaviour but lose the "
        "incubation-compatible Active>Passive pattern.",
        "Bottom line: ORBm & BMAp uniquely combine (i) Active-selective rise+dip, (ii) Active≥Passive through "
        "Withdrawal (incubation), and (iii) canonical addiction relevance.",
    ],
    subtitle="Systematic scan confirms ORBm/BMAp are the strongest Active-selective, incubation-compatible regions",
)
slide_image_big(
    "Fig 9 · Cluster-1/4 regions with the ORBm/BMAp signature (Active-selective + Active≥Passive)",
    SCAN / "Fig09_withdrawal_dip_scan.png",
    "Red = already selected (ORBm, BMAp, LM, RE). Green = other regions with the same signature (GPi, VAL, PAR, "
    "APr, MA — non-canonical). Every OTHER same-shape region has a deeper dip that falls below Passive "
    "(behaviour-match).")

# ================= Passive-SELECTIVE withdrawal regions =================
slide_text(
    "Passive data are informative too — Passive-selective Withdrawal regions",
    [
        "New question: are there regions where Active WINS every phase EXCEPT Withdrawal, where the PASSIVE "
        "(yoked) group instead peaks? These would capture the passive group's distinct withdrawal state.",
        "We scanned all 4 universal clusters for: Active>Passive at During, Post & Reinstatement, but "
        "Passive>Active ONLY at Withdrawal. Result: 27 regions (mostly clusters 1 & 4; cluster 3 adds AVPV/AVP).",
        "Checking the RAW trajectories, the regions where the PASSIVE group genuinely PEAKS at Withdrawal "
        "(not merely Active dropping) are: IGL, SI, EPd, AVP, ProS.",
    ],
    subtitle="Active wins everywhere except Withdrawal, where Passive peaks — a withdrawal-specific Passive signal",
)
slide_image_big(
    "Fig 10 · Regions where Passive wins ONLY in Withdrawal (all other phases Active wins)",
    SCAN / "Fig10_passive_wins_withdrawal_only.png",
    "Heatmap of Active−Passive by phase for the 27 strict regions. The Withdrawal column (boxed) is all-blue "
    "(Passive wins); every other phase is red (Active wins). Labels coloured by cluster.")
slide_image_big(
    "Fig 10b · Raw trajectories — does the PASSIVE group actually peak at Withdrawal?",
    SCAN / "Fig10b_passive_withdrawal_trajectories.png",
    "Shaded = Withdrawal. Blue (Passive) rising above red (Active) at Withdrawal confirms a Passive-selective "
    "withdrawal peak (clear in IGL, SI, EPd, AVP; PVi/SUB flip is mostly Active dropping).")
slide_text(
    "Refinement — keep only regions that ALSO carry the Active craving trajectory",
    [
        "Among the Passive-selective set (IGL, SI, EPd, AVP, ProS), we kept only those where the ACTIVE group "
        "still shows the canonical craving shape: a PEAK at Post AND a RISE at Reinstatement (rebound after "
        "the withdrawal dip).",
        "IGL and ProS were dropped — their Active peak is at During, not Post.",
        "FINAL Passive-informative regions = SI, EPd, AVP: a DOUBLE DISSOCIATION in one region — the Active "
        "morphine trajectory (Post peak + Reinstatement rebound) AND a Passive withdrawal-specific peak.",
        "These are attractive because a single Xenium region can report BOTH the craving ensemble (Active) and "
        "the withdrawal-state ensemble (Passive).",
    ],
    subtitle="Passive peaks at Withdrawal AND Active shows Post-peak + Reinstatement-rise → SI, EPd, AVP",
)
slide_image_big(
    "Fig 11 · Final Passive-informative regions (double dissociation): SI, EPd, AVP",
    SCAN / "Fig11_passive_selective_final.png",
    "Raw density (cells/mm³). Active (red) peaks at Post and rebounds at Reinstatement; Passive (blue) peaks "
    "at Withdrawal (shaded). Both signals coexist in SI, EPd, AVP.")

# ---------------------------------------------------------------- Part 2 header
slide_text("Part 2 · Zoom-in on each selected region (group level)",
           ["For each of the 7 regions we plot BOTH density (cells/mm³) and cell count across the 4 phases, "
            "separately for Active and Passive.",
            "Design note: TRAP is a TERMINAL snapshot — each mouse contributes to ONE phase only. So we show "
            "every individual mouse as a dot and connect only the GROUP MEAN ± SEM across phases (different mice "
            "per phase). No per-mouse cross-phase spaghetti (that would be fabricated)."],
           subtitle="density + cell count · Active vs Passive · individual mice shown")

for reg in REGIONS:
    name, why = REGION_INFO[reg]
    cap = (f"{reg} = {name} (Cluster {REGION_CLUSTER[reg]}). Relevance: {why}. "
           "Dots = individual mice; lines = group mean ± SEM across (different) mice per phase.")
    slide_image_big(f"{reg} · density & cell count over phases  (Cluster {REGION_CLUSTER[reg]})",
                    ZOOM / f"{reg}_density_count_by_phase.png", cap)

# ---------------------------------------------------------------- Part 3 header
slide_text("Part 3 · Within-phase Active vs Passive (one bar = one mouse)",
           ["A fair, same-phase comparison: within each phase we place every individual Active mouse next to "
            "every individual Passive mouse (bars), with the group mean as a horizontal line.",
            "This is WITHIN each phase (During: Active mice vs Passive mice, etc.) — a valid comparison of "
            "same-phase animals. It is NOT a cross-phase line (no spaghetti), consistent with the terminal design.",
            "It exposes mouse-to-mouse variability that the group means hide (important with small n)."],
           subtitle="per-region, per-phase individual-mouse comparison · density + cell count")

for reg in REGIONS:
    name, why = REGION_INFO[reg]
    cap = (f"{reg} = {name} (Cluster {REGION_CLUSTER[reg]}). Each bar = one mouse; red = Active, blue = Passive; "
           "horizontal line = group mean. Comparison is within each phase.")
    slide_image_big(f"{reg} · within-phase Active vs Passive  (Cluster {REGION_CLUSTER[reg]})",
                    ZOOM / f"{reg}_withinphase_ActiveVsPassive.png", cap)

# ---------------------------------------------------------------- closing
slide_text("Summary & next steps",
           ["Unbiased clustering filtered 138 forebrain regions → 4 clusters → 2 behaviour-aligned clusters "
            "→ 7 Xenium candidate regions (BMAp, LM, RE, CP, ORBm, CA, AId).",
            "Zoom-in confirms the pattern at the region level in BOTH density and cell count: Active > Passive "
            "at Post & Reinstatement in the selected regions.",
            "Selection logic: we kept INCUBATION regions (Active>Passive even in Withdrawal). If exact behaviour "
            "matching is preferred, PVi, SI, ENT (etc.) are the alternative set — decision to confirm with Mark/Greg.",
            "Active-selectivity check: ORBm & BMAp are Active-selective (rise+dip in Active only, Passive flat) and "
            "stay Active≥Passive through Withdrawal → strongest next targets. AId moves in both groups; CA is "
            "passive-like — so ORBm/BMAp are prioritised.",
            "Passive-informative regions: SI, EPd, AVP show a DOUBLE DISSOCIATION — Active craving trajectory "
            "(Post peak + Reinstatement rise) plus a Passive withdrawal-specific peak — capturing both states in "
            "one region.",
            "All individual mice shown; comparisons respect the terminal design (no fabricated trajectories).",
            "Next: confirm 7 vs narrow to 2–3 regions with Mark/Greg; finalize Xenium gene/probe list; "
            "cell-type-informed GPCR manipulation."])

candidates = [PPTX] + [PPTX.with_name(f"{PPTX.stem}_v{i}.pptx") for i in range(2, 12)]
out = None
for cand in candidates:
    try:
        prs.save(str(cand))
        out = cand
        break
    except PermissionError:
        continue
if out is None:
    from datetime import datetime
    out = PPTX.with_name(f"{PPTX.stem}_{datetime.now():%H%M%S}.pptx")
    prs.save(str(out))
if out != PPTX:
    print(f"(earlier versions were locked/open) saved to: {out}")
print(f"Saved: {out}\nSlides: {len(prs.slides._sldIdLst)}")
