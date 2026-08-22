"""
Colleague explanation deck: biological aim → workflow → each final figure
with meaning + code path on the same slide.

Run:  python build_colleague_explain_pptx.py
Out:  TRAP_colleague_explain_figure_code.pptx (+ .pdf if PowerPoint available)
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

PACK = Path(__file__).resolve().parent
FIG = PACK / "figures"
OUT_PPTX = PACK / "TRAP_colleague_explain_figure_code.pptx"
OUT_PDF = PACK / "TRAP_colleague_explain_figure_code.pdf"

GH = "https://github.com/limserenahansol/TRAP_analysis"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xB3, 0x00, 0x00)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
BLUE = RGBColor(0x15, 0x65, 0xC0)
DARK = RGBColor(0x22, 0x22, 0x22)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def title(s, text, color=NAVY, size=24):
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(SW - 0.8), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    tb.text_frame.word_wrap = True


def footer_code(s, code_lines: list[str]):
    tb = s.shapes.add_textbox(Inches(0.4), Inches(SH - 0.85), Inches(SW - 0.8), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(code_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11)
        p.font.color.rgb = BLUE
        p.font.name = "Consolas"


def bullets(s, items, left, top, w, h, size=14, colors=None):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("•  " + it) if not it.startswith("•") else it
        p.font.size = Pt(size)
        p.font.color.rgb = (colors[i] if colors and i < len(colors) else DARK)
        p.space_after = Pt(6)


def img_fit(s, path: Path, left, top, maxw, maxh):
    if not path.exists():
        tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(maxw), Inches(0.4))
        tb.text_frame.paragraphs[0].text = f"[Missing: {path.name}]"
        return
    iw, ih = Image.open(path).size
    r = min(maxw / iw, maxh / ih)
    w, h = iw * r, ih * r
    s.shapes.add_picture(
        str(path),
        Inches(left + (maxw - w) / 2),
        Inches(top + (maxh - h) / 2),
        width=Inches(w),
        height=Inches(h),
    )


def slide_text(t, items, subtitle=None, code=None, colors=None):
    s = prs.slides.add_slide(BLANK)
    title(s, t)
    y = 1.0
    if subtitle:
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(SW - 1), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.italic = True
        p.font.color.rgb = GREY
        y = 1.35
    bullets(s, items, 0.6, y, SW - 1.2, SH - y - 1.0, size=16, colors=colors)
    if code:
        footer_code(s, code)
    return s


def slide_fig_explain(t, img: Path, meaning: list[str], code: list[str], img_w=8.4):
    s = prs.slides.add_slide(BLANK)
    title(s, t)
    img_fit(s, img, 0.3, 0.95, img_w, SH - 1.95)
    bullets(s, meaning, img_w + 0.5, 1.05, SW - img_w - 0.9, SH - 2.1, size=13)
    footer_code(s, code)
    return s


def slide_two_figs(t, img1: Path, img2: Path, meaning: list[str], code: list[str]):
    s = prs.slides.add_slide(BLANK)
    title(s, t)
    img_fit(s, img1, 0.25, 0.95, 6.0, 4.6)
    img_fit(s, img2, 6.4, 0.95, 6.5, 4.6)
    bullets(s, meaning, 0.4, 5.7, SW - 0.8, 0.9, size=12)
    footer_code(s, code)
    return s


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

s = prs.slides.add_slide(BLANK)
tb = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(SW - 1.6), Inches(3.2))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "TRAP whole-brain analysis — colleague explanation"
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = NAVY
for line, sz, col in [
    ("Figure + meaning + code for every step used in the final story", 18, GREY),
    ("Findings: A = ORBm / BMAp (Active ≫ Passive) · B = COAa (Passive-biased)", 16, DARK),
    (f"Code: {GH}", 14, BLUE),
    ("Built for AI-benchmark navigation (final-PDF figures only)", 14, GREY),
]:
    p2 = tf.add_paragraph()
    p2.text = line
    p2.font.size = Pt(sz)
    p2.font.color.rgb = col

slide_text(
    "1 · Biological aim",
    [
        "Primary: find regions with the largest, most consistent Active ≫ Passive TRAP density across morphine phases.",
        "Secondary: find Passive-biased regions (Passive > Active at Post + Withdrawal) that may track generalized motivation.",
        "Design: Active (contingent morphine SA) vs Passive (yoked) · During → Post → Withdrawal → Reinstatement.",
        "Readout: TRAP+ cell density (cells / sample volume mm³), L/R pooled · Allen atlas · calculated_mm3.",
        "Why it matters: Active-selective + Passive-flat → opioid-related seeking; Passive-biased → patient-relevant / general seeking candidate.",
    ],
    code=[
        "Data / cohort: TRAP_sample_manifest.csv · density workbook Hansol Lim 561 cell counts + density.xlsx",
        f"Repo: {GH}",
    ],
)

slide_text(
    "2 · Two findings (final story)",
    [
        "FINDING A — ORBm & BMAp: Active ≫ Passive; Passive flat → dual-role craving (opioid seeking + WD negative affect).",
        "FINDING B — COAa: Passive > Active at Post + Withdrawal → cue/context + WD state → generalized-motivation candidate.",
        "Why keep both: A isolates opioid-related Active seeking; B is the Passive / patient-relevant contrast.",
        "Why not only behavior-matched regions for A: those track both opioid craving and Passive generalized motivation.",
    ],
    colors=[RED, GREEN, BLUE, GREY],
)

slide_fig_explain(
    "3 · Framing — opioid seeking vs general seeking",
    FIG / "05_schematics" / "S03_schematic_three_circuits_clean.png",
    [
        "WHAT: schematic separating Active-selective craving from Passive / general seeking.",
        "WHY: Passive-flat at ORBm/BMAp argues against pure general-seeking accounts.",
        "COAa is the Passive-side / generalized candidate to manipulate separately.",
        "AI check: agent should state both stories, not only Active≫Passive.",
    ],
    [
        "Figure source: forMark/schematic_three_circuits_clean.png",
        "Related scripts: trap_forMark_figs.py · circuit justification PPT builders",
    ],
)

slide_text(
    "4 · Workflow funnel (why cluster first, then zoom in)",
    [
        "Problem: hundreds of atlas regions → testing one-by-one is under-powered and unstructured.",
        "Step 1 FILTER: forebrain_no_bs (~138) · universal PCA display + k-means K=4 on Active−Passive fingerprint across 4 phases.",
        "Step 2 KEEP: only clusters 1 & 4 match motivation timeline (A>P at Post & Reinstatement).",
        "Step 3 SHORTLIST: cluster4 BMAp/LM/RE/CP + cluster1 ORBm/CA/AId (= 7 regions).",
        "Step 4 ZOOM IN: density + cell count + mouse-level → only ORBm & BMAp survive as Finding A; Finding B from cluster 3 → COAa.",
    ],
    subtitle="~800+ atlas → hierarchy567 → forebrain_no_bs → K=4 → clusters 1&4 → 7 → ORBm/BMAp + COAa",
    code=[
        "MATLAB: RUN_PIPELINE_ALL.m · Step_03_region_clustering… · Step_13_universal_cluster_viz/…",
        "Python story figures: trap_region_selection_ppt.py · trap_region_zoomin.py · trap_build_pptx.py",
    ],
)

slide_fig_explain(
    "5 · Fig 1 — Clustering method + PCA map",
    FIG / "01_pipeline_clustering" / "S05_Fig01_clustering_method_PCAmap.png",
    [
        "WHAT: 138 forebrain regions in PC1–PC2, colored by K=4; selected 7 circled.",
        "Method: z-score region profile across samples → k-means (sq-Euclidean, 50 restarts).",
        "PCA is 2-D display only; clusters use the full profile.",
        "AI check: keep fixed Step-3 K=4 labels; silhouette/elbow are supporting only.",
    ],
    [
        "CODE: trap_region_selection_ppt.py  →  region_selection_PPT/Fig01_…",
        "Upstream MATLAB: Step_03 + Step_13 (roster 02_cluster_region_roster.csv)",
    ],
)

slide_fig_explain(
    "6 · Fig 3 — Why clusters 1 and 4",
    FIG / "01_pipeline_clustering" / "S06_Fig03_why_clusters_1_and_4.png",
    [
        "WHAT: per-phase bars of cluster-mean Active − Passive.",
        "Only clusters 1 & 4 show A>P at Post & Reinstatement (motivation timeline).",
        "Post ≈ Active morphine learning; WD ≈ Passive sucrose shift; Rein ≈ Active craving.",
        "AI check: select clusters by phase pattern, not silhouette alone.",
    ],
    [
        "CODE: trap_region_selection_ppt.py  →  Fig03_why_clusters_1_and_4.png",
        "Inputs: cluster_AP_split/Cluster{N}_region_AP_direction.csv",
    ],
)

slide_fig_explain(
    "7 · Fig 4 — Cluster 4 shortlist (BMAp, LM, RE, CP)",
    FIG / "01_pipeline_clustering" / "S07_Fig04_cluster4_region_selection.png",
    [
        "WHAT: cluster-4 regions × phase heatmap of Active − Passive.",
        "Boxed / selected: BMAp, LM, RE, CP (A>P in craving-relevant phases).",
        "BMAp later survives zoom-in as Finding A (with ORBm).",
        "AI check: recover BMAp from cluster-4 AP shortlist.",
    ],
    [
        "CODE: trap_region_selection_ppt.py  →  Fig04_cluster4_region_selection*.png",
        "MATLAB support: shared/trap_cluster_split_by_AP_direction.m (Step 13)",
    ],
)

slide_fig_explain(
    "8 · Fig 5 — Cluster 1 shortlist (ORBm, CA, AId)",
    FIG / "01_pipeline_clustering" / "S08_Fig05_cluster1_region_selection.png",
    [
        "WHAT: cluster-1 regions ranked / heatmapped by Active − Passive.",
        "Selected shortlist: ORBm, CA, AId.",
        "ORBm later survives zoom-in as Finding A (with BMAp).",
        "AI check: recover ORBm from cluster-1 shortlist.",
    ],
    [
        "CODE: trap_region_selection_ppt.py  →  Fig05_cluster1_region_selection*.png",
        "MATLAB support: shared/trap_cluster_split_by_AP_direction.m (Step 13)",
    ],
)

slide_text(
    "9 · Part 2 — Zoom-in logic (Finding A filter)",
    [
        "After the 7-region shortlist, inspect density AND cell count across phases, Active vs Passive, every mouse.",
        "Keep regions where Active rises at Post, dips at WD but stays ≫ Passive, and Passive stays flat.",
        "Drop regions where Passive moves with Active (shared context) or one mouse drives the mean.",
        "Survivors for Finding A: ORBm + BMAp only.",
    ],
    code=[
        "CODE: trap_region_zoomin.py  →  region_zoomin/<REG>_density_count_by_phase.png",
        "                →  region_zoomin/<REG>_withinphase_ActiveVsPassive.png",
    ],
)

slide_fig_explain(
    "10 · Finding A schematic — dual-role craving",
    FIG / "05_schematics" / "S10_schematic_findingA_dual_role_incubation.png",
    [
        "WHAT: conceptual dual-mode craving at ORBm/BMAp.",
        "Post: opioid seeking (Active high). WD: negative-affect craving persists while Passive flat.",
        "WHY shown: frames why Active≫Passive through WD is incubation-compatible.",
    ],
    [
        "Figure: forMark/schematic_findingA_dual_role_incubation.png",
        "Related: trap_findingA_bridge_fig.py · trap_forMark_figs.py",
    ],
)

slide_fig_explain(
    "11 · Finding A evidence — ORBm & BMAp raw density",
    FIG / "02_findingA_ORBm_BMAp" / "S11_evidence_active_ORBm_BMAp.png",
    [
        "WHAT: Active vs Passive density trajectories for ORBm & BMAp.",
        "Pattern: Active peaks Post, dips WD, rebounds Rein; Passive essentially flat.",
        "This is the Active-selective signature used to argue opioid-related seeking.",
    ],
    [
        "CODE: trap_forMark_figs.py / trap_hilo_ORBm_BMAp_COAa.py",
        "Out: forMark/evidence_active_ORBm_BMAp.png",
    ],
)

slide_fig_explain(
    "12 · Why not the other original-7 regions?",
    FIG / "02_findingA_ORBm_BMAp" / "S12_deprioritized_5regions_summary.png",
    [
        "WHAT: summary of why CA, AId, LM, RE, CP were deprioritized for Finding A.",
        "Typical failure modes: Passive not flat, both groups move together, or unstable mouse leverage.",
        "AI check: agent should not treat all 7 as equal final targets.",
    ],
    [
        "CODE / out: forMark/deprioritized_5regions_summary.png (forMark helpers)",
    ],
)

slide_fig_explain(
    "13 · ORBm zoom-in — density & cell count",
    FIG / "02_findingA_ORBm_BMAp" / "S13_ORBm_density_count_by_phase.png",
    [
        "WHAT: ORBm group-level density and cell count by phase (Active vs Passive).",
        "Confirms the Active-selective rise / WD dip / Rein rebound pattern.",
        "Density metric = calculated_mm3 (cells / sample volume).",
    ],
    [
        "CODE: trap_region_zoomin.py",
        "Out: region_zoomin/ORBm_density_count_by_phase.png",
    ],
)

slide_two_figs(
    "14 · ORBm mouse-level + paired slope",
    FIG / "02_findingA_ORBm_BMAp" / "S17_ORBm_withinphase_ActiveVsPassive.png",
    FIG / "02_findingA_ORBm_BMAp" / "S17_ORBm_density_paired_slope.png",
    [
        "Left: one bar = one mouse within-phase Active vs Passive. Right: paired-slope consistency across mice.",
        "Shows the effect is not a single-mouse artifact.",
    ],
    [
        "CODE: trap_region_zoomin.py · forMark/paired_slope/ORBm_density_paired_slope.png",
    ],
)

slide_fig_explain(
    "15 · BMAp zoom-in — density & cell count",
    FIG / "02_findingA_ORBm_BMAp" / "S18_BMAp_density_count_by_phase.png",
    [
        "WHAT: BMAp group-level density and cell count by phase.",
        "Same Active-selective / Passive-flat motif as ORBm → second Finding A region.",
        "Anatomically: basomedial amygdala posterior (emotional valence / addiction motivation).",
    ],
    [
        "CODE: trap_region_zoomin.py",
        "Out: region_zoomin/BMAp_density_count_by_phase.png",
    ],
)

slide_two_figs(
    "16 · BMAp mouse-level + paired slope",
    FIG / "02_findingA_ORBm_BMAp" / "S22_BMAp_withinphase_ActiveVsPassive.png",
    FIG / "02_findingA_ORBm_BMAp" / "S22_BMAp_density_paired_slope.png",
    [
        "Mouse-level Active vs Passive bars + paired slopes confirm BMAp consistency across animals.",
    ],
    [
        "CODE: trap_region_zoomin.py · forMark/paired_slope/BMAp_density_paired_slope.png",
    ],
)

slide_text(
    "17 · Finding B start — why Passive matters",
    [
        "Passive is not only a control: noncontingent exposure still engages cue/context and WD state.",
        "Regions with Passive > Active at Post + Withdrawal are candidates for generalized (non-opioid-selective) motivation.",
        "Manipulation prediction: COAa OFF may affect broad seeking (e.g. sucrose); ORBm/BMAp OFF should spare that more.",
        "This is a testable dissociation, not a proven causal result from TRAP alone.",
    ],
    colors=[GREEN, GREEN, BLUE, GREY],
)

slide_fig_explain(
    "18 · Finding B schematic — Passive withdrawal / patient-relevant",
    FIG / "05_schematics" / "S25_schematic_findingB_passive_withdrawal.png",
    [
        "WHAT: schematic for Passive-biased / patient-relevant motif.",
        "Links Post+WD Passive elevation to cue/context + withdrawal state.",
        "Sets up cluster-3 screen for Finding B candidates.",
    ],
    [
        "Figure: forMark/schematic_findingB_passive_withdrawal.png",
        "Related: trap_findingB_cluster_map.py",
    ],
)

slide_fig_explain(
    "19 · Cluster 3 heatmap — Finding B candidates",
    FIG / "03_findingB_COAa" / "S26_Fig_cluster3_FindingB_heatmap.png",
    [
        "WHAT: cluster-3 region heatmap / ranking for Passive-biased pattern.",
        "COAa chosen as best anatomical fit (LGv/FC can match numerically).",
        "AI check: agent must produce a second Passive-biased story, not only Finding A.",
    ],
    [
        "CODE: trap_fig_cluster3_heatmap.py · findingB_COAa_revision/",
        "Out: Fig_cluster3_FindingB_heatmap.png",
    ],
)

slide_fig_explain(
    "20 · COAa zoom-in — density & cell count",
    FIG / "03_findingB_COAa" / "S27_COAa_density_count_by_phase.png",
    [
        "WHAT: COAa density and cell count by phase, Active vs Passive.",
        "Pattern: Passive > Active especially at Post and Withdrawal.",
        "This is Finding B’s primary quantitative panel.",
    ],
    [
        "CODE: COAa family / P>A helpers (see trap_build_pptx_COAa_family.py)",
        "Out: forMark/COAa_family/COAa_density_count_by_phase.png",
    ],
)

slide_two_figs(
    "21 · COAa mouse-level + paired slope",
    FIG / "03_findingB_COAa" / "S28_COAa_withinphase_ActiveVsPassive.png",
    FIG / "03_findingB_COAa" / "S28_COAa_density_paired_slope.png",
    [
        "Mouse-level and paired-slope panels for COAa Passive > Active consistency.",
    ],
    [
        "CODE: P_gt_A_PostWD helpers · forMark/paired_slope/COAa_density_paired_slope.png",
    ],
)

slide_text(
    "22 · Code map (where to click on GitHub)",
    [
        "Pipeline drivers: init_TRAP_pipeline.m · trap_config.m · RUN_PIPELINE_ALL.m",
        "Clustering: Step_03_region_clustering_PCA_kmeans/ · Step_13_universal_cluster_viz/",
        "Final-story Python (repo root): trap_region_selection_ppt.py · trap_region_zoomin.py · trap_build_pptx.py",
        "Finding helpers: trap_fig_cluster3_heatmap.py · trap_hilo_ORBm_BMAp_COAa.py · trap_forMark_figs.py",
        "Pack offline copies: colleague_AI_benchmark_pack/code/",
        "Canonical outputs: TRAP_OUTPUT_calculated_mm3/region_selection_PPT/ · region_zoomin/ · forMark/",
    ],
    code=[GH],
)

slide_text(
    "23 · AI-benchmark rubric (what success looks like)",
    [
        "Recover biological aim: Active≫Passive primary + Passive-biased secondary.",
        "Reproduce funnel: forebrain → K=4 → keep clusters 1&4 → shortlist 7 → zoom survivors ORBm/BMAp.",
        "Name Finding A = ORBm + BMAp with Passive-flat Active-selective rationale.",
        "Name Finding B = COAa (Passive > Active at Post+WD); mention LGv/FC as numeric peers.",
        "Point to the correct scripts above for each figure — not the entire Step-13 PNG dump.",
        "Do not treat histology slides as pipeline-generated plots (they were inserted into the PPT).",
    ],
)

slide_text(
    "24 · Summary",
    [
        "Aim: map opioid-selective craving circuits vs generalized / Passive-side motivation.",
        "Method: cluster first (structure) → shortlist → zoom-in (individual regions + mice).",
        "Finding A: ORBm & BMAp — Active ≫ Passive, Passive flat.",
        "Finding B: COAa — Passive > Active at Post + Withdrawal.",
        "This deck = figure + meaning + code for the final story only.",
        "Companion docs in pack: 00_START_HERE.md · HOW_FINAL_PDF_WAS_MADE.md · FINAL_PDF_FIGURE_MAP.md",
    ],
    colors=[NAVY, NAVY, RED, GREEN, BLUE, GREY],
    code=[
        f"GitHub: {GH}",
        "Local pack: Research_Projects/02_TRAP_wholebrain/colleague_AI_benchmark_pack/",
    ],
)

prs.save(OUT_PPTX)
print("Wrote", OUT_PPTX)


def try_export_pdf(pptx: Path, pdf: Path) -> bool:
    try:
        import win32com.client  # type: ignore
    except Exception as e:
        print("No win32com; skip PDF export:", e)
        return False
    ppt = None
    try:
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        ppt.Visible = 1
        # 32 = ppSaveAsPDF
        presentation = ppt.Presentations.Open(str(pptx), WithWindow=False)
        presentation.SaveAs(str(pdf), 32)
        presentation.Close()
        print("Wrote", pdf)
        return True
    except Exception as e:
        print("PDF export failed:", e)
        return False
    finally:
        if ppt is not None:
            try:
                ppt.Quit()
            except Exception:
                pass


try_export_pdf(OUT_PPTX, OUT_PDF)
